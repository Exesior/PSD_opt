# -*- coding: utf-8 -*-
"""Write changelog to Excel file."""

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from datetime import datetime

# Excel file path
excel_path = r"C:\Studium\Master\MA\Code\Änderungsnachverfolgung.xlsx"

# Headers
headers = ["Zeitstempel", "Welche Änderung", "Welches Script (Pfad + Zeile)", "Wofür diese Änderung", "Warum diese Syntax"]

# Data rows
data = [
    [
        "2026-06-12 15:00",
        "Neues Array `liquid_volume` initialisiert",
        "mcpbe/src/mcpbe/mcpbe_base.py, Zeile ~418",
        "Neues Flüssigkeitsvolumen-Attribut pro Partikel anlegen",
        "np.zeros(self._cap, dtype=float) – Array mit Nullen der Größe Capacity, Gleitkommazahlen"
    ],
    [
        "2026-06-12 15:00",
        "liquid_new bei Capacity-Wachstum",
        "mcpbe/src/mcpbe/mcpbe_base.py, Zeile ~476-478",
        "Bei Verdopplung der Capacity muss liquid_volume synchron mitwachsen",
        "liquid_new[:self.a_tot] = self.liquid_volume[:self.a_tot] – Alte Werte in neues Array kopieren"
    ],
    [
        "2026-06-12 15:00",
        "liquid_dup bei Control-Volume Verdopplung",
        "mcpbe/src/mcpbe/mcpbe_base.py, Zeile ~520",
        "Bei Verdopplung des Kontrollvolumens werden alle Partikel dupliziert",
        "np.concatenate(...) – Array verdoppeln für CV-Duplikation"
    ],
    [
        "2026-06-12 15:00",
        "liquid_new bei neuer Capacity nach CV-Verdopplung",
        "mcpbe/src/mcpbe/mcpbe_base.py, Zeile ~529-532",
        "Auch bei neuer Allokation muss liquid_volume verdoppelt geschrieben werden",
        "Gleiche Logik wie V_new, X_new"
    ],
    [
        "2026-06-12 15:00",
        "liquid_dup in else-Zweig",
        "mcpbe/src/mcpbe/mcpbe_base.py, Zeile ~536",
        "Bei ausreichender Capacity: direkt in existierendes Array schreiben",
        "self.liquid_volume[:self.a_tot] = liquid_dup – Slice-Zuweisung"
    ],
]

# Create workbook
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Änderungsnachverfolgung"

# Style definitions
header_font = Font(bold=True, color="FFFFFF")
header_fill = openpyxl.styles.PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

# Write headers
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border

# Write data rows
for row_idx, row_data in enumerate(data, 2):
    for col_idx, value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.border = thin_border

# Set column widths
ws.column_dimensions['A'].width = 20
ws.column_dimensions['B'].width = 40
ws.column_dimensions['C'].width = 50
ws.column_dimensions['D'].width = 60
ws.column_dimensions['E'].width = 80

# Save
wb.save(excel_path)
print(f"Excel file saved to: {excel_path}")